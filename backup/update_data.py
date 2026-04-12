"""
update_data.py
==============
使用方法：
  1. 把新的 Excel 替换到 data/data_explorer.xlsx
  2. 在项目根目录运行：python update_data.py
  3. 脚本自动更新 data/explorer_data.json 并将数据内嵌进 data_explorer.html

依赖：pip install openpyxl
"""
import openpyxl, json, re, os, sys

BASE_DIR   = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, 'data', 'data_explorer.xlsx')
JSON_PATH  = os.path.join(BASE_DIR, 'data', 'explorer_data.json')
HTML_PATH  = os.path.join(BASE_DIR, 'data_explorer.html')

METHOD_RULES = {
    "微调": "微调", "增加输入通道": "输入通道",
    "特征融合": "特征融合", "与主干融合": "特征融合",
    "设计CNN": "特征融合", "后处理": "后处理校正",
    "图神经网络": "图GNN", "条件注入": "条件注入",
}
PRIO_MAP = {"🔴 必需": "red", "🟡 推荐": "yellow", "🟢 可选": "green"}
CAT_MAP = {
    "气象与环境": "met", "相邻/上游测风站数据": "met",
    "空间与地理地形": "geo", "风电场与机组状态": "farm",
    "风速本体及时序特征": "feat", "本体外扩展": "ext",
}
PAPER_FIELDS = [
    "paper_id","doi","uses_era5","title","year","open_source",
    "spatial_res","temporal_res","wind_height","terrain",
    "met_input","sensor_input","dataset_method","dataset_link",
    "spatial_input","local_met_input","temporal_feature",
    "output_form","algorithm","physics_fusion","loss_func",
    "performance","business_map",
]

def clean(v):
    return str(v or "").replace("\n"," ").replace("\r"," ").strip()

def parse_cats(raw):
    s = clean(raw).replace("+","；")
    parts = re.split(r"[；;]", s)
    cats = []
    for p in parts:
        p = p.strip()
        for kw, code in CAT_MAP.items():
            if kw in p and code not in cats:
                cats.append(code)
    return cats if cats else ["met"]

def parse_diff(s):
    return min(max(str(s).count("⭐"), 1), 4)

def extract_methods(lines):
    found = set()
    for line in lines:
        for kw, label in METHOD_RULES.items():
            if kw in line:
                found.add(label)
    return sorted(found)

def extract_paper_ids(text):
    return [int(x) for x in re.findall(r"[\[{](\d+)[\]}]", str(text))]

def parse_excel(path):
    print(f"读取 Excel: {path}")
    wb = openpyxl.load_workbook(path)

    # Sheet1
    ws1 = wb["ERA5和其他数据源组合调研(补充算法组的数据源信息)"]
    sheet1 = {}
    for row in ws1.iter_rows(min_row=2, max_row=ws1.max_row, min_col=1, max_col=24, values_only=True):
        bid = clean(row[1])
        if not bid:
            continue
        sheet1[bid] = {
            "id": bid, "cat_label": clean(row[0]), "usage_flag": clean(row[2]),
            "name": clean(row[3]), "cat_raw": row[4], "onto": clean(row[5]),
            "res": clean(row[6]), "prio_raw": clean(row[7]), "fmt": clean(row[8]),
            "avail": clean(row[9]), "source": clean(row[10]), "diff_raw": clean(row[16]),
            "cost": clean(row[18]), "delay": clean(row[19]), "history": clean(row[20]),
            "geo": clean(row[21]), "necessity": clean(row[22]), "effect": clean(row[23]),
        }
    print(f"  Sheet1: {len(sheet1)} 条")

    # Sheet2 — 合并单元格续行修复
    ws2 = wb["数据源对应的文献以及数据使用方法和条目需求"]
    sheet2 = {}
    last_bid = None
    for row in ws2.iter_rows(min_row=2, max_row=ws2.max_row, min_col=1, max_col=6, values_only=True):
        bid = clean(row[1])
        if bid and not bid.startswith("NEW-"):
            last_bid = bid
        current_bid = bid if (bid and not bid.startswith("NEW-")) else last_bid
        if not current_bid:
            continue
        usage_e = clean(row[4])
        paper_f = clean(row[5])
        if current_bid not in sheet2:
            sheet2[current_bid] = {"method_lines": [], "paper_req_raw": "", "paper_ids": set()}
        if usage_e:
            sheet2[current_bid]["method_lines"].append(usage_e)
        if paper_f:
            sheet2[current_bid]["paper_req_raw"] += paper_f + " "
            for pid in extract_paper_ids(paper_f):
                sheet2[current_bid]["paper_ids"].add(pid)
    print(f"  Sheet2: {len(sheet2)} 个编号")

    # Sheet3
    ws3 = wb["文献列表"]
    papers = {}
    for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row, min_col=1, max_col=23, values_only=True):
        pid = row[0]
        if not pid:
            continue
        try:
            pid_int = int(pid)
        except (ValueError, TypeError):
            continue
        papers[pid_int] = {PAPER_FIELDS[i]: clean(row[i]) for i in range(23)}
        papers[pid_int]["paper_id"] = pid_int
    print(f"  Sheet3: {len(papers)} 篇文献")

    return sheet1, sheet2, papers

def merge(sheet1, sheet2, papers):
    items = []
    for bid, s1 in sheet1.items():
        s2 = sheet2.get(bid, {})
        ml = s2.get("method_lines", [])
        items.append({
            "id": bid, "name": s1["name"],
            "cat": parse_cats(s1["cat_raw"]),
            "cat_label": s1["cat_label"], "onto": s1["onto"],
            "prio": PRIO_MAP.get(s1["prio_raw"], "green"),
            "fmt": s1["fmt"], "res": s1["res"],
            "avail": s1["avail"], "source": s1["source"],
            "cost": s1["cost"], "delay": s1["delay"],
            "history": s1["history"], "geo": s1["geo"],
            "necessity": s1["necessity"], "effect": s1["effect"],
            "diff": parse_diff(s1["diff_raw"]),
            "method": extract_methods(ml),
            "method_lines": ml,
            "paper_req": s2.get("paper_req_raw", "").strip(),
            "paper_ids": sorted(s2.get("paper_ids", set())),
            "usage_flag": s1["usage_flag"],
        })
    print(f"  合并完成: {len(items)} 条 items")
    return items, list(papers.values())

def write_json(payload, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    print(f"JSON 已写入: {path} ({os.path.getsize(path)/1024:.1f} KB)")

def embed_html(payload, path):
    if not os.path.exists(path):
        print(f"HTML 不存在，跳过: {path}")
        return
    js_data = json.dumps(payload, ensure_ascii=False, separators=(",",":"))
    if "\n" in js_data or "\r" in js_data:
        print("ERROR: js_data 含真实换行符，中止")
        sys.exit(1)
    with open(path, encoding="utf-8") as f:
        html = f.read()
    m_start = "const EXPLORER_DATA = "
    m_end   = ";\n    let ITEMS"
    if m_start not in html or m_end not in html:
        print("ERROR: HTML 锚点未找到，请检查文件")
        sys.exit(1)
    i1 = html.index(m_start) + len(m_start)
    i2 = html.index(m_end)
    with open(path, "w", encoding="utf-8") as f:
        f.write(html[:i1] + js_data + html[i2:])
    print(f"HTML 已更新: {path} ({os.path.getsize(path)/1024:.1f} KB)")

def main():
    print("=" * 48)
    print("  data_explorer 数据更新脚本")
    print("=" * 48)
    if not os.path.exists(EXCEL_PATH):
        print(f"ERROR: Excel 不存在: {EXCEL_PATH}")
        sys.exit(1)
    sheet1, sheet2, papers = parse_excel(EXCEL_PATH)
    print("\n合并数据...")
    items, papers_list = merge(sheet1, sheet2, papers)
    payload = {"items": items, "papers": papers_list}
    print("\n写入文件...")
    write_json(payload, JSON_PATH)
    embed_html(payload, HTML_PATH)
    print("\n完成！重新打开 data_explorer.html 即可看到最新数据。")

if __name__ == "__main__":
    main()
